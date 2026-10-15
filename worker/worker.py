#!/usr/bin/env python3
"""One-request extraction/model boundary. Caller authorizes paths before invocation."""
import contextlib
import json
import os
from pathlib import Path
import sys
import subprocess
import tempfile

TEXT_EXTENSIONS = {'.txt', '.md', '.rst', '.csv', '.tsv', '.json', '.jsonl', '.xml', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.log', '.py', '.go', '.rs', '.java', '.js', '.jsx', '.ts', '.tsx', '.c', '.h', '.cpp', '.hpp', '.cs', '.sh', '.bash', '.sql', '.dfy', '.css', '.scss', '.vue', '.svelte', '.rb', '.php', '.swift', '.kt', '.kts', '.r', '.tex'}
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.tif', '.tiff', '.heic', '.heif'}
# CLIP sees 224 pixels; a small upright copy decodes faster and bounds memory.
EMBED_IMAGE_SIZE = (448, 448)


def ocr_pdf_page(path, number):
    """Rasterize only one bounded page, then recognize it with one CPU thread."""
    environment = dict(os.environ, OMP_THREAD_LIMIT='1', OMP_NUM_THREADS='1')
    with tempfile.TemporaryDirectory(prefix='nas-ocr-') as directory:
        output = str(Path(directory) / 'page')
        commands = [
            ['pdftoppm', '-f', str(number), '-l', str(number), '-singlefile',
             '-scale-to', '1600', '-png', str(path), output],
            ['tesseract', output + '.png', 'stdout', '--psm', '3', '-l', 'eng'],
        ]
        text = ''
        for command in commands:
            try:
                completed = subprocess.run(command, capture_output=True, text=True,
                                           timeout=90, env=environment)
            except (OSError, subprocess.TimeoutExpired) as exc:
                raise ValueError('PDF OCR unavailable or timed out') from exc
            if completed.returncode != 0:
                raise ValueError('PDF OCR failed')
            text = completed.stdout
        return text


def text_quality_suspect(text):
    """Conservative glyph-layer warning, not a claim that prose is understood."""
    import unicodedata
    characters = [character for character in text if not character.isspace()]
    if len(characters) < 20:
        return False
    readable = sum(character.isalnum() for character in characters)
    malformed = sum(character == '\ufffd' or unicodedata.category(character) in {'Co', 'Cc', 'Cs', 'Cn'} for character in characters)
    import itertools
    letters = sum(character.isalpha() for character in characters)
    punctuation = sum(unicodedata.category(character).startswith('P') for character in characters)
    repeated_punctuation = any(unicodedata.category(character).startswith('P') and sum(1 for _ in group) >= 8
                               for character, group in itertools.groupby(text))
    dense_glyph_token = any(len(token) >= 20 and sum(unicodedata.category(character).startswith('P') for character in token) * 2 > len(token) for token in text.split())
    glyph_pattern = (repeated_punctuation or dense_glyph_token) and letters * 4 < len(characters) and punctuation * 2 >= len(characters)
    return readable * 4 < len(characters) or malformed * 5 >= len(characters) or glyph_pattern


def usable_ocr_replacement(original, candidate):
    if not candidate.strip():
        return False
    if not original.strip():
        return True
    if text_quality_suspect(candidate):
        return False
    # Reject tiny OCR fragments replacing a substantial existing layer.
    readable = sum(character.isalnum() for character in candidate)
    original_readable = sum(character.isalnum() for character in original)
    return readable >= min(20, max(3, original_readable))


def extract(request, document_window=False):
    path = Path(request['path'])
    force_ocr = request.get('force_ocr', False)
    if type(force_ocr) is not bool:
        raise ValueError('force_ocr must be a boolean')
    if force_ocr and path.suffix.lower() != '.pdf':
        raise ValueError('force_ocr is supported only for PDF files')
    ocr_enabled = (os.environ.get('NAS_OCR') or 'true').lower() == 'true'
    if force_ocr and not ocr_enabled:
        raise ValueError('force_ocr requires NAS_OCR=true')
    max_bytes = int(request.get('max_bytes', 32 * 1024 * 1024))
    max_chars = int(request.get('max_chars', 200000))
    offset = request.get('offset', 0) if document_window else 0
    if max_bytes <= 0 or max_chars <= 0:
        raise ValueError('max_bytes and max_chars must be positive')
    if path.stat().st_size > max_bytes:
        raise ValueError('file exceeds max_bytes')
    result = {'kind': 'document', 'sections': [], 'metadata': {}, 'truncated': False}
    remaining = max_chars
    skip = offset
    section_seen = False
    def add(location, text, continuation=False):
        nonlocal remaining, skip, section_seen
        text = str(text)
        if not text:
            return remaining > 0
        if document_window:
            if section_seen and not continuation:
                text = '\n' + text
            section_seen = True
            consumed = min(skip, len(text))
            skip -= consumed
            text = text[consumed:]
        if len(text) > remaining:
            result['truncated'] = True
        bounded = text[:remaining]
        if bounded:
            if document_window and result['sections'] and result['sections'][-1]['location'] == location:
                result['sections'][-1]['text'] += bounded
            else:
                result['sections'].append({'location': location, 'text': bounded})
        remaining -= len(bounded)
        return remaining > 0
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        import codecs
        decoder = codecs.getincrementaldecoder('utf-8')('strict')
        text_parts = []
        characters = 0
        with path.open('rb') as stream:
            while True:
                raw = stream.read(65536)
                if any(b < 32 and b not in (9, 10, 13) for b in raw):
                    raise ValueError('binary file disguised as text')
                try:
                    piece = decoder.decode(raw, final=not raw)
                except UnicodeDecodeError as exc:
                    raise ValueError('text file must contain valid UTF-8') from exc
                if document_window:
                    add('text', piece, continuation=True)
                elif characters < max_chars:
                    text_parts.append(piece[:max_chars - characters])
                characters += len(piece)
                if not raw:
                    break
        if not document_window:
            add('text', ''.join(text_parts))
            result['truncated'] = characters > max_chars
    elif suffix in {'.html', '.htm'}:
        from html.parser import HTMLParser
        class Parser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.hidden = 0
            def handle_starttag(self, tag, attrs):
                if tag in {'script', 'style'}:
                    self.hidden += 1
            def handle_endtag(self, tag):
                if tag in {'script', 'style'}:
                    self.hidden = max(0, self.hidden - 1)
            def handle_data(self, data):
                if not self.hidden:
                    add('html', data)
        parser = Parser()
        with path.open('r', encoding='utf-8', errors='replace') as stream:
            while remaining:
                piece = stream.read(4096)
                if not piece:
                    break
                parser.feed(piece)
        if remaining == 0:
            result['truncated'] = True
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        result['metadata']['pages'] = len(reader.pages)
        result['metadata'].update(needs_ocr=False, ocr_pages=[], ocr_used=False, ocr_assessment_complete=False,
                                  text_quality_suspect=False, suspect_pages=[], ocr_forced=force_ocr,
                                  ocr_failed_pages=[], ocr_replaced_pages=[])
        for number, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ''
            except Exception:
                if not force_ocr:
                    raise
                text = ''
            suspect = text_quality_suspect(text)
            if suspect:
                result['metadata']['text_quality_suspect'] = True
                result['metadata']['suspect_pages'].append(number)
            if force_ocr or suspect or not text.strip():
                if ocr_enabled:
                    try:
                        candidate = ocr_pdf_page(path, number)
                    except ValueError:
                        # Automatic recovery must not discard an existing layer.
                        if force_ocr or not text.strip():
                            raise
                        result['metadata']['ocr_failed_pages'].append(number)
                    else:
                        result['metadata']['ocr_pages'].append(number)
                        result['metadata']['ocr_used'] = True
                        if force_ocr or usable_ocr_replacement(text, candidate):
                            text = candidate
                            result['metadata']['ocr_replaced_pages'].append(number)
                if not text.strip() or text_quality_suspect(text):
                    result['metadata']['needs_ocr'] = True
            result['metadata']['ocr_assessment_complete'] = number == len(reader.pages)
            if not add(f'page:{number}', text):
                result['truncated'] |= number < len(reader.pages)
                break
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        for index, block in enumerate(doc.iter_inner_content(), 1):
            text = block.text if hasattr(block, 'text') else '\n'.join('\t'.join(cell.text for cell in row.cells) for row in block.rows)
            if not add(f'block:{index}', text):
                break
    elif suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if not add(f'sheet:{sheet.title}:row:{number}', '\t'.join('' if value is None else str(value) for value in row)):
                        break
                if not remaining:
                    break
        finally:
            book.close()
    elif suffix == '.pptx':
        from pptx import Presentation
        for number, slide in enumerate(Presentation(path).slides, 1):
            for shape in slide.shapes:
                if getattr(shape, 'has_text_frame', False):
                    add(f'slide:{number}', shape.text)
                elif getattr(shape, 'has_table', False):
                    for row in shape.table.rows:
                        add(f'slide:{number}', '\t'.join(cell.text for cell in row.cells))
                if not remaining:
                    break
            if not remaining:
                break
    elif suffix in IMAGE_EXTENSIONS:
        from PIL import Image, ExifTags
        with open_image(path) as image:
            if image.width * image.height > Image.MAX_IMAGE_PIXELS:
                raise ValueError('image exceeds pixel limit')
            result['kind'] = 'photo'
            if document_window:
                result['metadata']['text_extraction_supported'] = False
            result['metadata'].update(width=image.width, height=image.height, format=image.format)
            exif = image.getexif()
            for key, value in exif.items():
                name = ExifTags.TAGS.get(key, str(key))
                if name in {'DateTime', 'DateTimeOriginal', 'Make', 'Model', 'Software', 'ImageDescription'}:
                    result['metadata'][name] = str(value)[:1024]
    else:
        raise ValueError(f'unsupported file type: {suffix or "no extension"}')
    return result


def open_image(path):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 25_000_000
    if Path(path).suffix.lower() in {'.heic', '.heif'}:
        import pillow_heif
        pillow_heif.register_heif_opener()
    return Image.open(path)


def embedding_image(path):
    from PIL import ImageOps
    with open_image(path) as image:
        # JPEG decoders can shrink while decoding; other formats ignore this hint.
        image.draft('RGB', EMBED_IMAGE_SIZE)
        upright = ImageOps.exif_transpose(image).convert('RGB')
    upright.thumbnail(EMBED_IMAGE_SIZE)
    return upright


def read_document(request):
    offset = request.get('offset', 0)
    limit = request.get('max_chars', 12000)
    if type(offset) is not int or offset < 0:
        raise ValueError('offset must be a nonnegative integer')
    if type(limit) is not int or not 1 <= limit <= 20000:
        raise ValueError('max_chars must be an integer from 1 to 20000')
    # One extra Unicode character distinguishes exact EOF from another page.
    window = dict(request, offset=offset, max_chars=limit + 1)
    result = extract(window, document_window=True)
    length = sum(len(section['text']) for section in result['sections'])
    complete = length <= limit
    if not complete:
        result['sections'][-1]['text'] = result['sections'][-1]['text'][:-1]
        if not result['sections'][-1]['text']:
            result['sections'].pop()
    text = ''.join(section['text'] for section in result['sections'])
    result.pop('truncated')
    result.update(text=text, offset=offset, complete=complete,
                  next_offset=None if complete else offset + len(text))
    result['metadata']['extracted_stream'] = 'nonempty sections joined by newline; plain text unchanged'
    return result


def load_embedding_classes():
    from fastembed import TextEmbedding, ImageEmbedding
    return TextEmbedding, ImageEmbedding


class ModelCache:
    """Retain at most one model; process shutdown releases its runtime memory."""
    def __init__(self):
        self.key = None
        self.model = None

    def get(self, cls, model_name):
        if self.key != model_name:
            import gc
            self.model = None
            self.key = None
            gc.collect()
            self.model = cls(model_name=model_name, threads=1,
                             providers=['CPUExecutionProvider'],
                             cache_dir=str(Path(os.environ.get('NAS_DATA', '/data')) / 'models'))
            self.key = model_name
        return self.model


def handle(request, model_cache=None):
    operation = request.get('operation', request.get('op'))
    if operation == 'extract':
        return extract(request)
    if operation == 'read_document':
        return read_document(request)
    if operation not in {'embed', 'image_embed', 'image_query'}:
        raise ValueError('unsupported operation')
    values = [request['path']] if operation == 'image_embed' else request['texts']
    if not isinstance(values, list) or not values or len(values) > 32 or any(not isinstance(value, str) for value in values):
        raise ValueError('embedding input must contain 1 to 32 strings')
    if operation != 'image_embed' and sum(len(value) for value in values) > 200000:
        raise ValueError('embedding input exceeds character limit')
    if operation == 'image_embed':
        # Check decoding dimensions and input size before loading heavyweight models.
        extract({'path': request['path'], 'max_bytes': request.get('max_bytes', 32 * 1024 * 1024), 'max_chars': 1})
        values = [embedding_image(request['path'])]
    TextEmbedding, ImageEmbedding = load_embedding_classes()
    model_name = {'embed': 'BAAI/bge-small-en-v1.5', 'image_embed': 'Qdrant/clip-ViT-B-32-vision', 'image_query': 'Qdrant/clip-ViT-B-32-text'}[operation]
    cls = ImageEmbedding if operation == 'image_embed' else TextEmbedding
    model = (model_cache if model_cache is not None else ModelCache()).get(cls, model_name)
    vectors = [vector.tolist() if hasattr(vector, 'tolist') else list(vector) for vector in model.embed(values, batch_size=1)]
    dimension = 384 if operation == 'embed' else 512
    if len(vectors) != len(values) or any(len(vector) != dimension for vector in vectors):
        raise ValueError('unexpected embedding dimensions')
    return {'vectors': vectors, 'dimension': dimension, 'model': model_name}


def serve():
    cache = ModelCache()
    for line in sys.stdin:
        try:
            with contextlib.redirect_stdout(sys.stderr):
                result = handle(json.loads(line), cache)
                response = json.dumps(result, allow_nan=False)
        except Exception as error:
            response = json.dumps({'error': str(error)})
        print(response, flush=True)
    return 0


def main():
    if sys.argv[1:] == ['--serve']:
        return serve()
    try:
        request = json.load(sys.stdin)
        with contextlib.redirect_stdout(sys.stderr):
            result = handle(request)
        print(json.dumps(result, allow_nan=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
