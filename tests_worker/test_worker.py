import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('extraction_worker', ROOT / 'worker/worker.py')
worker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(worker)


@pytest.fixture(autouse=True)
def disable_external_ocr(monkeypatch):
    monkeypatch.setenv('NAS_OCR', 'false')


def extract(path, **kw):
    return worker.handle({'operation': 'extract', 'path': str(path), 'max_bytes': 100000, 'max_chars': 1000, **kw})


def test_text_bounds(tmp_path):
    p = tmp_path / 'code.py'
    p.write_text('hello world')
    result = extract(p, max_chars=5)
    assert result['sections'] == [{'location': 'text', 'text': 'hello'}]
    assert result['truncated'] is True
    assert result['kind'] == 'document'
    with pytest.raises(ValueError, match='max_bytes'):
        extract(p, max_bytes=2)


def test_binary_disguised_as_text_is_rejected(tmp_path):
    p = tmp_path / 'project.cfg'
    for content in [b'A\x00binary', b'ABC\x01\x02binary', b'\xff\xfeinvalid']:
        p.write_bytes(content)
        with pytest.raises(ValueError, match='binary|UTF-8'):
            extract(p, max_chars=1)


def test_unicode_text_is_preserved(tmp_path):
    p = tmp_path / 'notes.txt'
    p.write_text('Music 🎵 中文\n', encoding='utf-8')
    assert extract(p)['sections'][0]['text'] == 'Music 🎵 中文\n'


def test_docx(tmp_path):
    p = tmp_path / 'test.docx'
    doc = Document()
    doc.add_paragraph('Alpha')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Beta'
    doc.save(p)
    assert 'Alpha' in str(extract(p)['sections'])
    assert 'Beta' in str(extract(p)['sections'])


def test_pdf_page_location(tmp_path):
    p = tmp_path / 'test.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(p)
    result = extract(p)
    assert result['kind'] == 'document'
    assert result['metadata']['pages'] == 1
    assert result['metadata']['needs_ocr'] is True


def test_photo_and_unsupported(tmp_path):
    p = tmp_path / 'test.png'
    Image.new('RGB', (5, 7)).save(p)
    result = extract(p)
    assert result['kind'] == 'photo'
    assert result['metadata']['width'] == 5
    assert result['sections'] == []
    p = tmp_path / 'backup.zip'
    p.write_bytes(b'x')
    with pytest.raises(ValueError, match='unsupported'):
        extract(p)


def test_embedding_models_and_dimensions(monkeypatch, tmp_path):
    image_path = tmp_path / 'model.jpg'
    Image.new('RGB', (5, 7)).save(image_path)
    calls = []
    class FakeModel:
        def __init__(self, **kwargs):
            calls.append(kwargs)
            self.dimension = 512 if 'clip' in kwargs['model_name'] else 384
        def embed(self, values, **kwargs):
            calls.append((list(values), kwargs))
            return [[0.1] * self.dimension]
    monkeypatch.setattr(worker, 'load_embedding_classes', lambda: (FakeModel, FakeModel))
    for operation, payload, dimension in [('embed', {'texts': ['hello']}, 384), ('image_query', {'texts': ['cat']}, 512), ('image_embed', {'path': str(image_path)}, 512)]:
        result = worker.handle({'operation': operation, **payload})
        assert len(result['vectors'][0]) == dimension
    assert all(c['threads'] == 1 for c in calls if isinstance(c, dict))
    assert all(c['providers'] == ['CPUExecutionProvider'] for c in calls if isinstance(c, dict))


def test_json_protocol_error():
    result = subprocess.run([sys.executable, str(ROOT / 'worker/worker.py')], input=json.dumps({'operation': 'bogus'}), capture_output=True, text=True)
    assert result.returncode == 1
    assert 'error' in json.loads(result.stdout)


def test_spreadsheet_and_presentation(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    p = tmp_path / 'book.xlsx'
    book = Workbook()
    book.active.append(['Heading', 42])
    book.save(p)
    sections = extract(p)['sections']
    assert sections == [{'location': 'sheet:Sheet:row:1', 'text': 'Heading\t42'}]
    p = tmp_path / 'slides.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = 'Slide content'
    deck.save(p)
    assert extract(p)['sections'] == [{'location': 'slide:1', 'text': 'Slide content'}]


def test_html_hides_scripts_and_byte_limit(tmp_path):
    p = tmp_path / 'page.html'
    p.write_text('<p>Visible</p><script>secret()</script><style>hidden</style>')
    result = extract(p)
    assert result['sections'] == [{'location': 'html', 'text': 'Visible'}]
    with pytest.raises(ValueError, match='positive'):
        extract(p, max_chars=0)


def test_embedding_input_limit(monkeypatch):
    monkeypatch.setattr(worker, 'load_embedding_classes', lambda: (_ for _ in ()).throw(AssertionError('models must not load for invalid request')))
    with pytest.raises(ValueError, match='1 to 32'):
        worker.handle({'operation': 'embed', 'texts': []})
    with pytest.raises(ValueError, match='character limit'):
        worker.handle({'operation': 'embed', 'texts': ['x' * 200001]})


def read_document(path, offset=0, max_chars=20000):
    return worker.handle({'operation': 'read_document', 'path': str(path), 'offset': offset, 'max_chars': max_chars})


def test_document_pages_entire_large_unicode_text(tmp_path):
    p = tmp_path / 'large.txt'
    content = ('🎵中文abc\n' * 35000)
    p.write_text(content, encoding='utf-8')
    chunks = []
    offset = 0
    while True:
        result = read_document(p, offset, 20000)
        assert result['offset'] == offset
        assert len(result['text']) <= 20000
        assert result['text'] == ''.join(s['text'] for s in result['sections'])
        chunks.append(result['text'])
        if result['complete']:
            assert result['next_offset'] is None
            break
        assert result['next_offset'] == offset + len(result['text'])
        offset = result['next_offset']
    assert ''.join(chunks) == content
    assert len(content) > 200000
    assert read_document(p, len(content), 1)['complete'] is True


def test_document_bounds_and_exact_eof(tmp_path):
    p = tmp_path / 'small.txt'
    p.write_text('abc🎵')
    result = read_document(p, 0, 4)
    assert result['text'] == 'abc🎵'
    assert result['complete'] is True
    assert read_document(p, 500, 4)['text'] == ''
    for offset, limit in [(-1, 5), (0, 0), (0, 20001), (True, 5), (0, 1.5)]:
        with pytest.raises(ValueError, match='offset|max_chars'):
            read_document(p, offset, limit)


def test_document_docx_section_offsets(tmp_path):
    p = tmp_path / 'whole.docx'
    doc = Document()
    doc.add_paragraph('Alpha🎵')
    doc.add_paragraph('Beta')
    doc.add_paragraph('Gamma')
    doc.save(p)
    full = 'Alpha🎵\nBeta\nGamma'
    pages = [read_document(p, offset, 4) for offset in range(0, len(full), 4)]
    assert ''.join(page['text'] for page in pages) == full
    assert read_document(p, 7, 4)['sections'] == [{'location': 'block:2', 'text': 'Beta'}]


def test_document_pdf_page_offsets_and_scanned_metadata(monkeypatch, tmp_path):
    import pypdf
    p = tmp_path / 'whole.pdf'
    p.write_bytes(b'fake')
    class Page:
        def __init__(self, text): self.text = text
        def extract_text(self): return self.text
    class Reader:
        def __init__(self, path): self.pages = [Page('First'), Page('Second'), Page('')]
    monkeypatch.setattr(pypdf, 'PdfReader', Reader)
    result = read_document(p, 4, 20)
    assert result['text'] == 't\nSecond'
    assert [s['location'] for s in result['sections']] == ['page:1', 'page:2']
    assert result['metadata']['needs_ocr'] is True
    assert result['complete'] is True
    monkeypatch.setattr(pypdf, 'PdfReader', lambda path: type('Reader', (), {'pages': [Page(''), Page('')]})())
    result = read_document(p)
    assert result['text'] == '' and result['complete'] is True
    assert result['metadata']['needs_ocr'] is True


def test_document_photo_does_not_claim_text_extraction(tmp_path):
    p = tmp_path / 'photo.jpg'
    Image.new('RGB', (5, 7)).save(p)
    result = read_document(p)
    assert result['text'] == '' and result['kind'] == 'photo'
    assert result['metadata']['text_extraction_supported'] is False


def test_scanned_pdf_ocr_shared_with_index_and_read(monkeypatch, tmp_path):
    import pypdf
    p = tmp_path / 'scan.pdf'
    p.write_bytes(b'fake')
    class Page:
        def __init__(self, text): self.text = text
        def extract_text(self): return self.text
    monkeypatch.setattr(pypdf, 'PdfReader', lambda path: type('Reader', (), {'pages': [Page('Embedded'), Page('')]})())
    calls = []
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda path, number: calls.append(number) or 'Scanned text')
    monkeypatch.setenv('NAS_OCR', 'true')
    for result in [extract(p), read_document(p)]:
        assert result['metadata']['ocr_pages'] == [2]
        assert result['metadata']['ocr_used'] is True
        assert result['metadata']['needs_ocr'] is False
        assert 'Scanned text' in str(result['sections'])
    assert calls == [2, 2]
    monkeypatch.setenv('NAS_OCR', 'false')
    result = read_document(p)
    assert result['metadata']['needs_ocr'] is True
    assert calls == [2, 2]


def test_pdf_ocr_commands_bound_raster_and_threads(monkeypatch, tmp_path):
    calls = []
    def run(command, **kwargs):
        calls.append((command, kwargs))
        return subprocess.CompletedProcess(command, 0, stdout='OCR content' if command[0] == 'tesseract' else '', stderr='')
    monkeypatch.setattr(worker.subprocess, 'run', run)
    assert worker.ocr_pdf_page(tmp_path / 'scan.pdf', 7) == 'OCR content'
    raster, recognizer = calls
    assert raster[0][raster[0].index('-scale-to')+1] == '1600'
    assert raster[0][raster[0].index('-f')+1] == '7'
    assert raster[0][raster[0].index('-l')+1] == '7'
    assert '-singlefile' in raster[0]
    assert recognizer[0][-4:] == ['--psm', '3', '-l', 'eng']
    assert all(call[1]['timeout'] == 90 for call in calls)
    assert all(call[1]['env']['OMP_THREAD_LIMIT'] == '1' for call in calls)


def test_pdf_ocr_failure_is_explicit(monkeypatch, tmp_path):
    monkeypatch.setattr(worker.subprocess, 'run', lambda command, **kwargs: subprocess.CompletedProcess(command, 1, stdout='', stderr='private document content'))
    with pytest.raises(ValueError, match='OCR') as error:
        worker.ocr_pdf_page(tmp_path / 'scan.pdf', 1)
    assert 'private document content' not in str(error.value)


def test_real_pdf_document_sections_and_eof(tmp_path):
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    p = tmp_path / 'pages.pdf'
    writer = PdfWriter()
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    for text in ['First page', 'Second page']:
        page = writer.add_blank_page(width=200, height=200)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 12 Tf 10 100 Td (' + text + ') Tj ET').encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(p)
    full = 'First page\nSecond page'
    chunks = [read_document(p, offset, 5) for offset in range(0, len(full), 5)]
    assert ''.join(chunk['text'] for chunk in chunks) == full
    assert chunks[-1]['complete'] is True
    assert read_document(p, 11, 11)['sections'] == [{'location': 'page:2', 'text': 'Second page'}]


def test_document_read_rejects_binary_after_return_window(tmp_path):
    p = tmp_path / 'late.txt'
    p.write_bytes(b'A' * 70000 + b'\x00binary')
    with pytest.raises(ValueError, match='binary'):
        read_document(p, 0, 1)


def fake_pdf(monkeypatch, tmp_path, texts):
    import pypdf
    path = tmp_path / 'glyphs.pdf'
    path.write_bytes(b'fake')
    pages = [type('Page', (), {'extract_text': lambda self, text=text: text})() for text in texts]
    monkeypatch.setattr(pypdf, 'PdfReader', lambda path: type('Reader', (), {'pages': pages})())
    return path


def test_force_ocr_replaces_existing_layer(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['Embedded placeholder'])
    calls = []
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda path, number: calls.append(number) or 'Complete recovered document text')
    result = worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})
    assert result['text'] == 'Complete recovered document text'
    assert result['metadata']['ocr_forced'] is True
    assert result['metadata']['ocr_pages'] == [1]
    assert calls == [1]


def test_force_ocr_validation(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['Normal prose'])
    with pytest.raises(ValueError, match='NAS_OCR'):
        worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})
    monkeypatch.setenv('NAS_OCR', 'true')
    for value in [1, 'true', None]:
        with pytest.raises(ValueError, match='boolean'):
            worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': value})
    path = tmp_path / 'text.txt'
    path.write_text('Text')
    with pytest.raises(ValueError, match='PDF'):
        worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})


@pytest.mark.parametrize('text', ['♠♣♦♥♩♪♫♬☀☁' * 11, '\ue001\ue002\ue003\ue004' * 30, '\ufffd' * 110])
def test_suspicious_layers_trigger_shared_ocr(monkeypatch, tmp_path, text):
    path = fake_pdf(monkeypatch, tmp_path, [text])
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda path, number: 'Recovered useful prose from the page.')
    for result in [extract(path), read_document(path)]:
        assert result['metadata']['text_quality_suspect'] is True
        assert result['metadata']['suspect_pages'] == [1]
        assert result['metadata']['ocr_forced'] is False
        assert result['metadata']['needs_ocr'] is False
        assert result['sections'][0]['text'] == 'Recovered useful prose from the page.'


@pytest.mark.parametrize('text', ['A normal document paragraph, with punctuation and 123 numbers.', '中文文档可以包含中文标点。继续阅读这些文字。', 'Музыкальный документ содержит нормальный русский текст.', 'Title', '42', '2026-10-11'])
def test_readable_layers_are_not_suspect(monkeypatch, tmp_path, text):
    path = fake_pdf(monkeypatch, tmp_path, [text])
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda *args: pytest.fail('Readable layer must not invoke OCR'))
    result = read_document(path)
    assert result['text'] == text
    assert result['metadata']['text_quality_suspect'] is False
    assert result['metadata']['suspect_pages'] == []


@pytest.mark.parametrize('replacement', ['', '♠♣♦♥' * 50])
def test_suspect_layer_preserved_when_ocr_empty_or_worse(monkeypatch, tmp_path, replacement):
    original = '♠♣♦♥' * 25 + 'Original123'
    path = fake_pdf(monkeypatch, tmp_path, [original])
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda *args: replacement)
    result = read_document(path)
    assert result['text'] == original
    assert result['metadata']['needs_ocr'] is True
    assert result['metadata']['text_quality_suspect'] is True


def test_suspect_layer_flagged_without_ocr(monkeypatch, tmp_path):
    text = '♠♣♦♥' * 30
    path = fake_pdf(monkeypatch, tmp_path, [text])
    result = read_document(path)
    assert result['text'] == text
    assert result['metadata']['needs_ocr'] is True
    assert result['metadata']['suspect_pages'] == [1]


def test_music_glyph_layer_with_repeated_punctuation(monkeypatch, tmp_path):
    # Same 110-character scale and low-letter/repeated-dot pattern as the failing PDF.
    text = 'Title\n' + '.' * 31 + '\n1234567890 ' + ',.-' * 15 + '\nHeading123456789'
    path = fake_pdf(monkeypatch, tmp_path, [text])
    result = read_document(path)
    assert result['metadata']['text_quality_suspect'] is True
    assert result['metadata']['suspect_pages'] == [1]


def test_prose_with_dot_leader_remains_readable(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['Chapter title and normal readable words .............. 42'])
    result = read_document(path)
    assert result['metadata']['text_quality_suspect'] is False


def test_force_ocr_is_explicit_override_even_short_or_empty(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['Long embedded text that the user deliberately overrides.'])
    monkeypatch.setenv('NAS_OCR', 'true')
    for candidate in ['Hi', '']:
        monkeypatch.setattr(worker, 'ocr_pdf_page', lambda *args: candidate)
        result = worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})
        assert result['text'] == candidate
        assert result['metadata']['needs_ocr'] is (candidate == '')


def test_empty_ocr_environment_uses_enabled_default(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['Embedded'])
    monkeypatch.setenv('NAS_OCR', '')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda *args: 'OCR')
    assert worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})['text'] == 'OCR'


def test_force_ocr_recovers_broken_embedded_extraction(monkeypatch, tmp_path):
    import pypdf
    path = tmp_path / 'broken.pdf'
    path.write_bytes(b'fake')
    class Page:
        def extract_text(self): raise ValueError('broken font mapping')
    monkeypatch.setattr(pypdf, 'PdfReader', lambda path: type('Reader', (), {'pages': [Page()]})())
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda *args: 'Recovered OCR')
    assert worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True})['text'] == 'Recovered OCR'


def test_force_ocr_multipage_pagination(monkeypatch, tmp_path):
    path = fake_pdf(monkeypatch, tmp_path, ['First layer', 'Second layer'])
    monkeypatch.setenv('NAS_OCR', 'true')
    monkeypatch.setattr(worker, 'ocr_pdf_page', lambda path, number: {1:'First recovered',2:'Second recovered'}[number])
    full = 'First recovered\nSecond recovered'
    pieces = []
    offset = 0
    while True:
        result = worker.handle({'operation': 'read_document', 'path': str(path), 'force_ocr': True, 'offset':offset, 'max_chars':5})
        assert result['metadata']['ocr_forced'] is True
        assert result['metadata']['ocr_pages'] == list(range(1,len(result['metadata']['ocr_pages'])+1))
        pieces.append(result['text'])
        if result['complete']: break
        offset=result['next_offset']
    assert ''.join(pieces) == full


def test_interleaved_music_glyph_punctuation_layer(monkeypatch, tmp_path):
    text = 'Title\n' + '.12..,3..4.,' * 5 + '\n42\nHeading\n' + ',.5.-6.' * 4
    path = fake_pdf(monkeypatch, tmp_path, [text])
    result = read_document(path)
    assert result['metadata']['text_quality_suspect'] is True


def test_automatic_ocr_failure_preserves_suspect_layer(monkeypatch, tmp_path):
    text = '♠♣♦♥' * 30
    path = fake_pdf(monkeypatch, tmp_path, [text])
    monkeypatch.setenv('NAS_OCR', 'true')
    def fail(*args): raise ValueError('PDF OCR unavailable')
    monkeypatch.setattr(worker, 'ocr_pdf_page', fail)
    result = read_document(path)
    assert result['text'] == text
    assert result['metadata']['ocr_failed_pages'] == [1]
    assert result['metadata']['needs_ocr'] is True


def test_iphone_heic_photos_are_indexed(tmp_path):
    import pillow_heif
    p = tmp_path / 'IMG_0001.HEIC'
    pillow_heif.from_pillow(Image.new('RGB', (8, 6), 'red')).save(p)
    result = extract(p)
    assert result['kind'] == 'photo'
    assert (result['metadata']['width'], result['metadata']['height']) == (8, 6)


def test_image_embedding_uses_small_upright_copy(monkeypatch, tmp_path):
    p = tmp_path / 'phone.jpg'
    exif = Image.Exif()
    exif[0x0112] = 6  # Stored sideways; viewers rotate it 90 degrees to portrait.
    Image.new('RGB', (1600, 800), 'blue').save(p, exif=exif)
    heic = tmp_path / 'phone.heic'
    import pillow_heif
    pillow_heif.from_pillow(Image.new('RGB', (1200, 900), 'green')).save(heic)
    received = []
    class FakeModel:
        def __init__(self, **kwargs):
            pass
        def embed(self, values, **kwargs):
            received.extend(values)
            return [[0.1] * 512 for _ in values]
    monkeypatch.setattr(worker, 'load_embedding_classes', lambda: (FakeModel, FakeModel))
    for path in (p, heic):
        worker.handle({'operation': 'image_embed', 'path': str(path)}, worker.ModelCache())
    portrait, landscape = received
    for image in received:
        assert isinstance(image, Image.Image) and image.mode == 'RGB'
        assert max(image.size) <= 448
    assert portrait.height > portrait.width
    assert landscape.width > landscape.height


def test_image_embedding_reports_photo_metadata(monkeypatch, tmp_path):
    p = tmp_path / 'beach.png'
    Image.new('RGB', (30, 20)).save(p)
    class FakeModel:
        def __init__(self, **kwargs):
            pass
        def embed(self, values, **kwargs):
            return [[0.1] * 512 for _ in values]
    monkeypatch.setattr(worker, 'load_embedding_classes', lambda: (FakeModel, FakeModel))
    result = worker.handle({'operation': 'image_embed', 'path': str(p)}, worker.ModelCache())
    assert result['kind'] == 'photo'
    assert (result['metadata']['width'], result['metadata']['height']) == (30, 20)
